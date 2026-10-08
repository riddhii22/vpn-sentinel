#!/usr/bin/env python3
"""SIH idea-template deck (6 slides) for PS 26160 — proposed full solution."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY2 = RGBColor(0x12, 0x33, 0x58)
ORANGE = RGBColor(0xE0, 0x5A, 0x26)
GOLD = RGBColor(0xF0, 0xA5, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF = RGBColor(0xF4, 0xF6, 0xF8)
INK = RGBColor(0x1A, 0x1C, 0x1E)
MUTED = RGBColor(0x4A, 0x55, 0x63)
GREEN = RGBColor(0x1F, 0x7A, 0x4C)
RED = RGBColor(0xB4, 0x23, 0x18)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD0, 0xD7, 0xDE)

W, H = Inches(13.333), Inches(7.5)


def _set_run(run, text, size=12, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def box(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line if line else fill
    sh.adjustments[0] = 0.06
    return sh


def rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line if line else fill
    return sh


def tb(slide, l, t, w, h, text, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    s = slide.shapes.add_textbox(l, t, w, h)
    tf = s.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    _set_run(p.add_run() if p.runs else p.runs[0] if False else None, "", size, bold, color, font)
    # python-pptx: first paragraph already has no run until we add
    p.clear() if hasattr(p, "clear") else None
    run = p.add_run()
    _set_run(run, text, size, bold, color, font)
    return s


def add_text(slide, l, t, w, h, lines, default_size=12, color=INK, align=PP_ALIGN.LEFT):
    """lines: list of str or (str, size, bold, color)."""
    s = slide.shapes.add_textbox(l, t, w, h)
    tf = s.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, str):
            txt, sz, b, c = item, default_size, False, color
        else:
            txt = item[0]
            sz = item[1] if len(item) > 1 else default_size
            b = item[2] if len(item) > 2 else False
            c = item[3] if len(item) > 3 else color
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(3)
        run = p.add_run()
        _set_run(run, txt, sz, b, c)
    return s


def footer(slide, n):
    rect(slide, 0, Inches(7.22), W, Inches(0.28), NAVY)
    add_text(slide, Inches(0.3), Inches(7.24), Inches(8), Inches(0.24),
             [("TEAM ENTROPY  ·  SIH 2026  ·  PS 26160  ·  NTRO — Cybersecurity", 10, False, WHITE)])
    add_text(slide, Inches(11.6), Inches(7.24), Inches(1.4), Inches(0.24),
             [(str(n), 11, True, ORANGE)], align=PP_ALIGN.RIGHT)


def header_bar(slide, title, num):
    rect(slide, 0, 0, W, Inches(0.92), NAVY)
    rect(slide, 0, Inches(0.92), W, Inches(0.08), ORANGE)
    add_text(slide, Inches(0.35), Inches(0.12), Inches(10), Inches(0.38),
             [(title, 22, True, WHITE)])
    add_text(slide, Inches(0.35), Inches(0.50), Inches(10), Inches(0.32),
             [("@SIH Idea submission template  ·  Software  ·  Cybersecurity", 11, False, RGBColor(0xC5, 0xD0, 0xDC))])
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(12.15), Inches(0.16), Inches(0.62), Inches(0.62))
    circ.fill.solid()
    circ.fill.fore_color.rgb = ORANGE
    circ.line.fill.background()
    add_text(slide, Inches(12.15), Inches(0.26), Inches(0.62), Inches(0.42),
             [(str(num), 18, True, WHITE)], align=PP_ALIGN.CENTER)


def slide1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, 0, Inches(0.18), H, ORANGE)
    add_text(s, Inches(0.55), Inches(0.35), Inches(12), Inches(0.35),
             [("SMART INDIA HACKATHON 2026", 14, True, ORANGE)])
    add_text(s, Inches(0.55), Inches(0.75), Inches(12), Inches(0.85),
             [("VPN SENTINEL", 40, True, WHITE)])
    add_text(s, Inches(0.55), Inches(1.55), Inches(12), Inches(0.45),
             [("AI-Powered IPsec VPN Protocol Analyzer and Security Assessment Framework", 18, False, RGBColor(0xE8, 0xEE, 0xF4))])

    fields = [
        ("Problem Statement ID", "26160"),
        ("Problem Statement Title", "AI-Powered IPsec VPN Protocol Analyzer and Security Assessment Framework"),
        ("Organization", "NTRO — Blockchain & Cybersecurity"),
        ("Theme", "Cybersecurity"),
        ("PS Category", "Software"),
        ("Team ID", "________________ (fill from SIH portal)"),
        ("Team Name", "ENTROPY"),
    ]
    y = Inches(2.2)
    for label, val in fields:
        box(s, Inches(0.55), y, Inches(3.3), Inches(0.48), NAVY2, RGBColor(0x2A, 0x4A, 0x70))
        add_text(s, Inches(0.65), y + Inches(0.08), Inches(3.1), Inches(0.35), [(label, 11, True, ORANGE)])
        box(s, Inches(3.95), y, Inches(8.7), Inches(0.48), RGBColor(0x15, 0x38, 0x60), RGBColor(0x2A, 0x4A, 0x70))
        add_text(s, Inches(4.1), y + Inches(0.08), Inches(8.4), Inches(0.35), [(val, 13, False, WHITE)])
        y += Inches(0.56)

    add_text(s, Inches(0.55), Inches(6.55), Inches(12), Inches(0.4),
             [("Lab testbed  ·  IKE/ESP/AH capture  ·  AI protocol ID  ·  Inner-traffic inference  ·  Security assessment  ·  Dashboard & reports", 12, False, RGBColor(0xB8, 0xC5, 0xD4))])
    add_text(s, Inches(0.55), Inches(7.05), Inches(12), Inches(0.28),
             [("1", 11, True, ORANGE)])


def slide2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFF)
    header_bar(s, "PROBLEM  ·  SOLUTION  ·  UNIQUENESS", 2)

    # Problem
    box(s, Inches(0.3), Inches(1.18), Inches(12.7), Inches(1.15), WHITE, LINE)
    add_text(s, Inches(0.45), Inches(1.22), Inches(12.4), Inches(0.28), [("PROBLEM (PS 26160)", 12, True, ORANGE)])
    add_text(s, Inches(0.45), Inches(1.50), Inches(12.4), Inches(0.75),
             [("IPsec VPNs look “secure” while still using IKEv1, weak ciphers (DES/3DES), short DH groups, no PFS, long lifetimes, or transport mode. Misconfigurations are found only by expert Wireshark work. NTRO needs an AI-assisted framework that can generate known VPN configs, capture IKE/ESP/AH, identify protocol parameters, infer inner traffic without decryption, score the config, and report with confidence.", 12, False, INK)])

    cols = [
        (Inches(0.3), "PROPOSED SOLUTION",
         "VPN Sentinel is an end-to-end lab + analyzer:\n"
         "• Controllable IPsec testbed (StrongSwan): Tunnel/Transport, AES-128/256 GCM & CBC+HMAC, multiple DH groups, PFS on/off, IPv4/IPv6, web/video/voice inner traffic.\n"
         "• Capture pipeline: IKE (UDP 500/4500), ESP, AH, plus labelled PCAPs per scenario.\n"
         "• AI/ML: (1) protocol identification from IKE SA — version, mode, encryption, integrity/auth, DH, SA lifetime; (2) inner traffic type from ESP size/timing only — never decrypt.\n"
         "• Security engine: crypto strength, compliance, PFS, replay, cipher suite, metadata exposure.\n"
         "• Outputs: 0–100 risk, threat matrix, AI confidence, executive + technical reports, interactive dashboard, training/testing dataset."),
        (Inches(6.75), "INNOVATION / UNIQUENESS",
         "• RFC-correct dual parsers: IKEv1 attribute list (RFC 2409) vs IKEv2 Transform Substructure (RFC 7296 §3.3) — not a shared byte hunt.\n"
         "• Config is ground truth for training only; inference reads packets, not the lab YAML.\n"
         "• Encrypted ESP is never opened; web/video/voice predicted from metadata (Draper-Gil–style features + IsolationForest / supervised head).\n"
         "• One operator loop: scenario matrix → pcap → identify → assess → report — built for NTRO-style review, not a commercial VPN box.\n"
         "• Explainable findings (rule IDs + evidence packets) plus an AI confidence score on the traffic/protocol heads."),
    ]
    for left, title, body in cols:
        box(s, left, Inches(2.48), Inches(6.25), Inches(4.55), WHITE, LINE)
        rect(s, left, Inches(2.48), Inches(6.25), Inches(0.42), NAVY)
        add_text(s, left + Inches(0.15), Inches(2.52), Inches(6.0), Inches(0.35), [(title, 13, True, WHITE)])
        add_text(s, left + Inches(0.15), Inches(3.00), Inches(5.95), Inches(3.9), [(body, 12, False, INK)])
    footer(s, 2)


def slide3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFF)
    header_bar(s, "TECHNICAL APPROACH", 3)

    steps = [
        ("1. TESTBED", "StrongSwan / Docker\nTunnel · Transport\nAES-128/256\nGCM / CBC+HMAC\nDH groups · PFS\nIPv4 / IPv6\nWeb · Video · Voice"),
        ("2. CAPTURE", "tcpdump / PCAPNG\nIKE UDP 500\nNAT-T UDP 4500\nESP (proto 50)\nAH (proto 51)\nScenario labels\nbeside each file"),
        ("3. AI IDENTIFY", "IKE v1 vs v2\nEnc / integ / DH\nMode · lifetime\nSA characteristics\nESP 5-tuple flows\nSize · IAT · duration\nTraffic class + conf."),
        ("4. ASSESS", "Crypto strength\nConfig compliance\nPFS · replay\nCipher suite\nMetadata exposure\nThreat matrix\n0–100 risk score"),
        ("5. DELIVER", "React dashboard\nExec + tech report\nAI confidence\nJSON / HTML / PDF\nDocs + demo video\nTrain/test dataset"),
    ]
    x = Inches(0.28)
    for title, body in steps:
        box(s, x, Inches(1.18), Inches(2.42), Inches(3.55), WHITE, LINE)
        rect(s, x, Inches(1.18), Inches(2.42), Inches(0.42), ORANGE if title.startswith("3") else NAVY)
        add_text(s, x + Inches(0.08), Inches(1.22), Inches(2.26), Inches(0.36), [(title, 12, True, WHITE)], align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.1), Inches(1.68), Inches(2.22), Inches(2.95), [(body, 11, False, INK)])
        x += Inches(2.58)

    add_text(s, Inches(0.3), Inches(4.82), Inches(12.7), Inches(0.28),
             [("STACK (software edition — no payload decryption)", 12, True, NAVY)])

    tech = [
        ("Testbed", "Linux, Docker Compose, strongSwan/swanctl, iperf / curl / ffmpeg inner traffic"),
        ("Capture", "tcpdump, PCAP/PCAPNG, Scapy inventory (IKE, ESP, AH, NAT-T)"),
        ("Protocol AI", "RFC 2409 + RFC 7296 parsers; optional learned head on transform/feature vectors"),
        ("Traffic AI", "Flow features → IsolationForest (anomaly) + supervised web/video/voice on labelled lab set"),
        ("Assessment", "YAML policy (NIST SP 800-77 style) + additive 0–100 risk, threat matrix"),
        ("Product", "FastAPI, React + Vite dashboard, HTML/PDF reports, GitHub dataset + docs"),
    ]
    y = Inches(5.12)
    for i, (k, v) in enumerate(tech):
        col = i % 2
        row = i // 2
        left = Inches(0.3) + (Inches(6.45) if col else 0)
        top = y + Inches(row * 0.58)
        box(s, left, top, Inches(6.3), Inches(0.52), WHITE, LINE)
        add_text(s, left + Inches(0.1), top + Inches(0.04), Inches(1.5), Inches(0.42), [(k, 11, True, ORANGE)])
        add_text(s, left + Inches(1.6), top + Inches(0.04), Inches(4.55), Inches(0.42), [(v, 11, False, INK)])
    footer(s, 3)


def slide4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFF)
    header_bar(s, "FEASIBILITY AND VIABILITY", 4)

    blocks = [
        (Inches(0.3), NAVY, "FEASIBILITY",
         "• Open-source IPsec (strongSwan) and packet tools (Scapy, tcpdump) are production-proven.\n"
         "• Public IKE samples plus a small Docker lab cover the PS config axes without buying appliances.\n"
         "• PCAP upload is a safe MVP; live sniff stays on the team’s lab interface only.\n"
         "• Assessment rules map to NIST SP 800-77 / RFC 9395 — not a black-box score.\n"
         "• ML uses observable metadata (size, timing, direction). No ESP decrypt, no illegal intercept."),
        (Inches(4.55), ORANGE, "CHALLENGES",
         "• Child SA / Quick Mode often encrypted → mode, PFS, replay may be hidden unless the lab logs them.\n"
         "• Inner-traffic classes overlap; accuracy needs a labelled testbed, not two Wireshark files.\n"
         "• IKEv1 vs IKEv2 encodings differ — a shared 0x80 0x01 hunt fails silently on IKEv2.\n"
         "• IPv6 + every cipher×DH×PFS combo is combinatorial; we ship a documented scenario matrix, not 100% of combinations on day one.\n"
         "• False positives on anomaly scores if baselines are too small."),
        (Inches(8.8), GREEN, "MITIGATION / PLAN",
         "• Dual RFC parsers + Wireshark ground truth on every sample.\n"
         "• Lab YAML stored next to each PCAP as labels (never used as the only inference input).\n"
         "• IsolationForest first (anomaly), then supervised web/video/voice on lab data with hold-out metrics and confidence.\n"
         "• Findings stay explainable: packet evidence + rule ID; ML is a separate finding type.\n"
         "• Phased build: parsers → rules → ESP ML → reports → testbed matrix expansion."),
    ]
    for left, color, title, body in blocks:
        box(s, left, Inches(1.18), Inches(4.1), Inches(4.35), WHITE, LINE)
        rect(s, left, Inches(1.18), Inches(4.1), Inches(0.42), color)
        add_text(s, left + Inches(0.12), Inches(1.22), Inches(3.85), Inches(0.36), [(title, 13, True, WHITE)])
        add_text(s, left + Inches(0.12), Inches(1.70), Inches(3.85), Inches(3.7), [(body, 11, False, INK)])

    box(s, Inches(0.3), Inches(5.65), Inches(12.7), Inches(1.35), WHITE, LINE)
    add_text(s, Inches(0.45), Inches(5.70), Inches(12.4), Inches(0.3), [("VIABILITY FOR SIH EXTERNAL ROUND", 12, True, NAVY)])
    add_text(s, Inches(0.45), Inches(6.02), Inches(12.4), Inches(0.9),
             [("Local demo is enough for evaluation (no cloud). Dataset = lab PCAPs + public IKE captures. Team can show: (A) handshake decode matching Wireshark, (B) distinct findings for weak vs strong configs, (C) ESP metadata ML with a confidence number, (D) dashboard + executive/technical report. Hardware buy is not required. Success depends on labelled scenarios and honest limits on encrypted Child SAs — not on decrypting the VPN.", 12, False, INK)])
    footer(s, 4)


def slide5(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFF)
    header_bar(s, "IMPACT AND BENEFITS", 5)

    users = [
        ("NTRO / CYBER UNITS", "Faster than manual IPsec review on captured tunnels; repeatable scoring."),
        ("SOC / BLUE TEAM", "IKEv1, weak DH, missing PFS and odd ESP shapes rise as prioritized findings."),
        ("VPN ADMINS", "Config compliance vs a written policy (AES, DH, lifetime) without opening payloads."),
        ("TRAINING LABS", "Scenario matrix teaches “what a bad VPN looks like on the wire.”"),
        ("AUDITORS", "Executive summary + technical evidence pack (packets, rule IDs, confidence)."),
        ("INDIA / AATMANIRBHAR", "Open stack (Python, strongSwan) instead of only foreign appliance GUIs."),
    ]
    positions = [(0.3, 1.18), (4.55, 1.18), (8.8, 1.18), (0.3, 3.05), (4.55, 3.05), (8.8, 3.05)]
    for (x, y), (title, body) in zip(positions, users):
        box(s, Inches(x), Inches(y), Inches(4.1), Inches(1.72), WHITE, LINE)
        rect(s, Inches(x), Inches(y), Inches(0.12), Inches(1.72), ORANGE)
        add_text(s, Inches(x + 0.28), Inches(y + 0.1), Inches(3.7), Inches(0.4), [(title, 13, True, NAVY)])
        add_text(s, Inches(x + 0.28), Inches(y + 0.52), Inches(3.7), Inches(1.1), [(body, 13, False, INK)])

    box(s, Inches(0.3), Inches(4.95), Inches(12.7), Inches(2.05), NAVY)
    add_text(s, Inches(0.5), Inches(5.05), Inches(12.3), Inches(0.32), [("STRATEGIC BENEFIT (ALIGNED TO THE PS)", 13, True, ORANGE)])
    add_text(s, Inches(0.5), Inches(5.40), Inches(12.3), Inches(1.45),
             [("Reduces dependence on expert-only packet reading for IPsec hygiene. Creates a labelled Indian-lab dataset of IKE/ESP under known configs — the missing piece for “AI protocol ID” and “traffic type inside ESP.” Does not replace a VPN gateway or attack others’ networks. Direct PS outcomes: testbed, capture, AI identification, inner-traffic prediction without decrypt, assessment (crypto, PFS, replay, suite strength, metadata), reports (risk, matrix, confidence), dashboard, documentation, demo, train/test set.", 13, False, WHITE)])
    footer(s, 5)


def slide6(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFF)
    header_bar(s, "RESEARCH, REFERENCES & PS TRACEABILITY", 6)

    add_text(s, Inches(0.35), Inches(1.15), Inches(6.2), Inches(0.3), [("STANDARDS & LITERATURE", 13, True, NAVY)])
    refs = (
        "• RFC 2409 — IKEv1; RFC 7296 §3.3 — IKEv2 Transform Substructure\n"
        "• RFC 4303 ESP; RFC 3948 NAT-T; RFC 9395 — IKEv1 deprecation\n"
        "• NIST SP 800-77r1 — Guide to IPsec VPNs (crypto/policy language)\n"
        "• Draper-Gil et al., ICISSP 2016 — time-related features for VPN traffic\n"
        "• Internet-wide IKE measurement / ike-scan-style transform enumeration\n"
        "• IsolationForest (Liu et al.) for unsupervised anomaly on flow stats"
    )
    box(s, Inches(0.3), Inches(1.48), Inches(6.35), Inches(2.55), WHITE, LINE)
    add_text(s, Inches(0.45), Inches(1.55), Inches(6.1), Inches(2.4), [(refs, 12, False, INK)])

    add_text(s, Inches(6.85), Inches(1.15), Inches(6.1), Inches(0.3), [("PS 26160 CHECKLIST (IN SCOPE)", 13, True, NAVY)])
    chk = (
        "✓ VPN testbed: Tunnel/Transport, AES-128/256, GCM/CBC+HMAC, DH, PFS, IPv4/IPv6, traffic types\n"
        "✓ Capture: IKE, ESP, AH, normal/labelled inner traffic\n"
        "✓ AI protocol ID: version, mode, enc, auth, KE, SA traits\n"
        "✓ AI inner type: web/video/voice from ESP metadata + confidence\n"
        "✓ Assessment: strength, compliance, lifetime, replay, PFS, suite, metadata\n"
        "✓ Reports + dashboard + dataset + demo video + docs"
    )
    box(s, Inches(6.8), Inches(1.48), Inches(6.2), Inches(2.55), WHITE, LINE)
    add_text(s, Inches(6.95), Inches(1.55), Inches(5.95), Inches(2.4), [(chk, 12, False, INK)])

    add_text(s, Inches(0.35), Inches(4.15), Inches(12.5), Inches(0.3), [("RISK MODEL (PROPOSED, REQUIREMENTS-STYLE)", 13, True, NAVY)])
    box(s, Inches(0.3), Inches(4.48), Inches(12.7), Inches(1.0), WHITE, LINE)
    add_text(s, Inches(0.45), Inches(4.55), Inches(12.4), Inches(0.85),
             [("VPN Sentinel Risk Score = min(100, Σ finding scores). Example weights: IKEv1 30 · weak enc 20 · weak DH 20 · PFS off 15 · transport 10 · long lifetime 10 · ESP/traffic anomaly from ML (separate ID). Bands: 0–29 Low · 30–59 Moderate · 60–79 High · 80–100 Critical. AI confidence is reported beside protocol/traffic predictions — it does not replace cryptographic facts.", 12, False, INK)])

    box(s, Inches(0.3), Inches(5.62), Inches(12.7), Inches(1.38), NAVY)
    add_text(s, Inches(0.45), Inches(5.70), Inches(12.4), Inches(0.28), [("ONE LINE FOR JUDGES", 12, True, ORANGE)])
    add_text(s, Inches(0.45), Inches(6.02), Inches(12.4), Inches(0.85),
             [("VPN Sentinel will generate known IPsec worlds, capture them, read what the handshake actually advertised, infer what the encrypted tunnel behaves like, and produce an explainable risk report — without ever decrypting ESP.", 14, False, WHITE)])
    footer(s, 6)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide1(prs)
    slide2(prs)
    slide3(prs)
    slide4(prs)
    slide5(prs)
    slide6(prs)
    out = "/workspace/docs/VPN_Sentinel_SIH_2026_ENTROPY_PS26160.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
