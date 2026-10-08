#!/usr/bin/env python3
"""Patch the team's attached SIH PPT (slides 2–5 only). Template chrome unchanged."""

from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

SRC = Path("/home/ubuntu/.cursor/projects/workspace/uploads/VPN_SENTINEL__1__2ce1.pptx")
UI = Path("/home/ubuntu/.cursor/projects/workspace/assets/f47f64be-2c75-4d94-bed0-689537d242c7.png")
CROP = Path("/workspace/docs/assets/vpn-sentinel-dashboard-slide2.png")
OUT = Path("/workspace/VPN_Sentinel_SIH2026_IDEA.pptx")
INK = RGBColor(0x00, 0x00, 0x00)


def crop_ui() -> None:
    im = Image.open(UI).convert("RGB")
    im.crop((0, 176, 1898, 1002)).save(CROP, "PNG")


def set_box(shape, text: str, size: int = 11, bold_first: bool = False) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bool(bold_first)
        run.font.name = "Arial"
        run.font.color.rgb = INK


def replace_picture(slide, name: str, image_path: Path) -> None:
    pic = next(sh for sh in slide.shapes if sh.name == name)
    left, top, width, height = pic.left, pic.top, pic.width, pic.height
    pic._element.getparent().remove(pic._element)
    slide.shapes.add_picture(str(image_path), left, top, width, height)


def main() -> None:
    crop_ui()
    shutil.copy(SRC, OUT)
    prs = Presentation(str(OUT))

    s2 = prs.slides[1]
    for sh in s2.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("• RFC-correct IKEv1"):
            set_box(
                sh,
                "• Dual RFC parsers: IKEv1 RFC 2409 TV attributes vs IKEv2 RFC 7296 §3.3 Transform Type/ID (ISAKMP version nibble >> 4).\n"
                "• NIST SP 800-77r1 cryptographic policy checks: cipher-suite verification, weak DH groups (1/2/5), hashing/integrity, IKEv1 deprecation (RFC 9395).\n"
                "• Zero-decryption behavioral profiling: IsolationForest on ESP flow metadata (size mean/variance, inter-arrival deltas, duration, byte volume). No payload inspection.\n"
                "• Aggregated security posture index: additive 0–100 risk score with finding IDs and an exportable HTML/JSON audit report.",
                10,
            )
        elif t.startswith("VPN Sentinel is an IPsec"):
            set_box(
                sh,
                "VPN Sentinel ingests operator-owned PCAPs, reconstructs IKE/ESP headers, and emits protocol facts plus a policy-backed risk score. Encrypted ESP payloads are never decrypted.",
                12,
            )
    replace_picture(s2, "Picture 3", CROP)

    s3 = prs.slides[2]
    for sh in s3.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("• FastAPI"):
            set_box(
                sh,
                "• FastAPI — asynchronous PCAP ingest and analysis API\n"
                "• Network processing: Scapy reconstruction of ISAKMP/IKE, ESP, AH, UDP/500, UDP/4500\n"
                "• parse_ikev1_sa — RFC 2409 short-form attributes (0x80 0x01 encryption TV)\n"
                "• parse_ikev2_sa — RFC 7296 §3.3 Transform Type/ID (no shared 0x80 0x01 hunt)\n"
                "• YAML policy — NIST-aligned crypto/config findings\n"
                "• ML core: IsolationForest on scaled ESP flow vectors (size/timing); no payload decryption; does not set the VPN risk score\n"
                "• Data interface: native PCAP/PCAPNG batch offline dumps\n"
                "• React (Vite) dashboard",
                10,
            )
        elif t.startswith("PCAP UPLOAD"):
            set_box(sh, "PACKET INGESTION\nFastAPI + Scapy (offline PCAP)", 10, True)
        elif t.startswith("IKEv1 PARSER"):
            set_box(sh, "IKEv1 PARSER\nVersion-nibble branch · RFC 2409", 10, True)
        elif t.startswith("IKEv2 PARSER"):
            set_box(sh, "IKEv2 PARSER\nVersion-nibble branch · RFC 7296", 10, True)
        elif "ISAKMP >> 4" in t:
            set_box(
                sh,
                "Processing splits on the IKE handshake: ISAKMP version nibble (byte 17 >> 4) selects the parser. Filename is ignored.",
                11,
            )
        elif t.startswith("RULE ENGINE"):
            set_box(sh, "NIST SP 800-77r1\nCryptographic Policy Check", 10, True)
        elif t.startswith("ML ENGINE"):
            set_box(sh, "IsolationForest\nTemporal & Structural Anomaly Processor", 9, True)
        elif t.startswith("RISK SCORE"):
            set_box(sh, "POSTURE INDEX  +  FINDINGS  →  DASHBOARD / AUDIT EXPORT", 11, True)
        elif t.startswith("GitHub:"):
            set_box(
                sh,
                "GitHub:  https://github.com/riddhii22/vpn-Sentinel-minor\n"
                "Repository status: parser core and NIST-aligned verification engine active (prototype testing). IsolationForest scores ESP metadata; it does not classify inner traffic.",
                12,
            )

    s4 = prs.slides[3]
    for sh in s4.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("• Offline PCAP analysis"):
            set_box(
                sh,
                "• Offline PCAP analysis removes live-capture operational and legal exposure.\n"
                "• Deterministic validation: parsers cross-checked against Wireshark-decoded IKEv1/IKEv2 captures (cipher, key length, DH, PSK/hash attributes).\n"
                "• Policy language tracks NIST SP 800-77r1 IPsec cryptographic guidance — not ad-hoc thresholds.",
                12,
            )
        elif t.startswith("• IKEv1 and IKEv2 need separate"):
            set_box(
                sh,
                "• IKEv1 and IKEv2 require separate SA encodings. A shared 0x80 0x01 hunt fails silently on IKEv2.\n"
                "• Encrypted-state traversal: Child SA / Quick Mode fields (PFS, tunnel vs transport, replay) that are ciphertext stay Unverified Configuration — the engine never invents values.\n"
                "• Payload encryption boundary: no inner-traffic class (YouTube/WhatsApp). Analysis stays on cryptographic strength and observable transmission metadata.",
                11,
            )
        elif t.startswith("• Binary anomaly detection"):
            set_box(
                sh,
                "• Non-invasive telemetry: IsolationForest profiles ESP structural flow signatures (size/timing). No DPI of ciphertext.\n"
                "• Lab testbed and labelled web/video/voice classes remain the expansion path under PS 26160.\n"
                "• Explainable intelligence: each alert maps to a rule ID plus packet evidence; ML is a separate telemetry score and does not overwrite cipher findings.",
                11,
            )

    s5 = prs.slides[4]
    for sh in s5.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("• Security teams:"):
            set_box(
                sh,
                "• Security teams: first-pass IPsec configuration audit without exclusive reliance on manual Wireshark trees.\n"
                "• Network / VPN admins: IKEv1, DES/3DES, and weak DH isolated with a recommendation.\n"
                "• Operational optimization: automates crypto-suite discovery and handshake compliance parsing.\n"
                "• Automated initial triage: surfaces high-risk handshake misconfigurations before any later lab-scale deployment.",
                12,
            )
        elif t.startswith("• Supports wider use"):
            set_box(
                sh,
                "• Supports correctly configured IPsec (IKEv2, modern ciphers, adequate DH) on enterprise and government paths.\n"
                "• Distinguishes “VPN is on” from “the handshake is cryptographically sound.”\n"
                "• Aligns with NTRO PS 26160 assessment and reporting: structural IKE/ESP parsing, policy auditing, and exportable compliance output. Testbed and inner-traffic classification remain the scale-out path.",
                12,
            )

    prs.save(str(OUT))
    shutil.copy(OUT, Path("/workspace/docs/VPN_Sentinel_SIH2026_IDEA.pptx"))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
