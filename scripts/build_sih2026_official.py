#!/usr/bin/env python3
"""Fill the official SIH 2026 IDEA template (colors/fonts/chrome unchanged)."""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

TEMPLATE = Path("/home/ubuntu/.cursor/projects/workspace/uploads/SIH2026-IDEA-Presentation-Format_15fb.pptx")
OUT = Path("/workspace/VPN_Sentinel_SIH2026_IDEA.pptx")

# Official template palette (theme + measured footer)
BLUE = RGBColor(0x00, 0x70, 0xC0)  # footer bar
DK = RGBColor(0x1F, 0x49, 0x7D)  # theme dk2
MID = RGBColor(0x4F, 0x81, 0xBD)
ORANGE = RGBColor(0xF7, 0x96, 0x46)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x00, 0x00, 0x00)
PALE = RGBColor(0xEE, 0xEC, 0xE1)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0x4F, 0x81, 0xBD)

TEAM = "ENTROPY"
TEAM_ID = "[TEAM ID — fill from SIH portal]"
GITHUB = "https://github.com/riddhii22/vpn-Sentinel-minor"
COMPLETION = "~20% of full PS  ·  parse + score + dashboard working"
HERO_SRC = Path("/workspace/docs/assets/vpn-sentinel-ui-capture.png")
HERO = Path("/workspace/docs/assets/vpn-sentinel-dashboard-slide2.png")


def set_run(run, text, size=12, bold=False, color=INK, font="Arial"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_marked_runs(paragraph, txt, sz, default_bold, color, font):
    """Render **this** as bold inside a sentence."""
    if "**" not in txt:
        run = paragraph.add_run()
        set_run(run, txt, sz, default_bold, color, font)
        return
    for i, part in enumerate(txt.split("**")):
        if part == "":
            continue
        run = paragraph.add_run()
        set_run(run, part, sz, (i % 2 == 1) or default_bold, color, font)


def fill_tf(shape, lines, default_size=14, color=INK, font="Arial", align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, str):
            txt, sz, b, c, font_i = item, default_size, False, color, font
        else:
            txt = item[0]
            sz = item[1] if len(item) > 1 else default_size
            b = item[2] if len(item) > 2 else False
            c = item[3] if len(item) > 3 else color
            font_i = item[4] if len(item) > 4 else font
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(3)
        p.level = 0
        add_marked_runs(p, txt, sz, b, c, font_i)


def prepare_dashboard_crop():
    """Drop browser chrome + taskbar; keep the real VPN Sentinel UI."""
    from PIL import Image

    HERO.parent.mkdir(parents=True, exist_ok=True)
    im = Image.open(HERO_SRC).convert("RGB")
    # (left, top, right, bottom) — dashboard only
    crop = im.crop((0, 176, 1898, 1002))
    crop.save(HERO, "PNG")


def delete_instruction_boxes(slide):
    """Remove original template prompt text boxes so they cannot overlap content."""
    markers = (
        "Proposed Solution (Describe",
        "Technologies to be used",
        "Analysis of the feasibility",
        "Potential impact on the target",
        "Details / Links of the reference",
    )
    sp_tree = slide.shapes._spTree
    for sh in list(slide.shapes):
        if not sh.has_text_frame:
            continue
        blob = sh.text_frame.text
        if any(m in blob for m in markers):
            sp_tree.remove(sh._element)


def card(slide, l, t, w, h, title, body, header=DK):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.adjustments[0] = 0.04
    sh.fill.solid()
    sh.fill.fore_color.rgb = WHITE
    sh.line.color.rgb = header
    sh.line.width = Pt(1.25)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, Inches(0.36))
    bar.fill.solid()
    bar.fill.fore_color.rgb = header
    bar.line.color.rgb = header
    tb = slide.shapes.add_textbox(l + Inches(0.08), t + Inches(0.02), w - Inches(0.12), Inches(0.32))
    fill_tf(tb, [(title, 13, True, WHITE, "Arial")])
    body_box = slide.shapes.add_textbox(l + Inches(0.1), t + Inches(0.40), w - Inches(0.18), h - Inches(0.48))
    lines = [(ln, 12, False, INK, "Arial") for ln in body.split("\n") if ln != ""]
    fill_tf(body_box, lines)
    return sh


def tag_box(slide, l, t, w, h, text, fill):
    """Helix-style coloured callout next to the product image."""
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.adjustments[0] = 0.18
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = fill
    tb = slide.shapes.add_textbox(l + Inches(0.04), t + Inches(0.04), w - Inches(0.08), h - Inches(0.06))
    fill_tf(tb, [(text, 10, True, WHITE, "Arial")], align=PP_ALIGN.CENTER)
    return sh


def arrow_box(slide, l, t, w, h, text, fill=DK):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = fill
    tb = slide.shapes.add_textbox(l + Inches(0.05), t + Inches(0.05), w - Inches(0.1), h - Inches(0.08))
    fill_tf(tb, [(text, 11, True, WHITE, "Arial")], align=PP_ALIGN.CENTER)
    return sh


def chevron(slide, l, t):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, l, t, Inches(0.28), Inches(0.22))
    sh.fill.solid()
    sh.fill.fore_color.rgb = ORANGE
    sh.line.color.rgb = ORANGE
    return sh


def set_oval_team(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and "Your Team Name" in sh.text_frame.text:
            fill_tf(sh, [(TEAM, 11, True, WHITE, "Arial")], align=PP_ALIGN.CENTER)
            try:
                sh.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            except Exception:
                pass


def delete_last_slide(prs):
    sldIdLst = prs.slides._sldIdLst
    sldId = list(sldIdLst)[-1]
    rId = sldId.get(qn("r:id"))
    sldIdLst.remove(sldId)
    try:
        prs.part.drop_rel(rId)
    except Exception:
        pass


def fill_slide1(slide):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.name == "TextBox 9":
            fill_tf(
                sh,
                [
                    ("Problem Statement ID –  26160", 20, True, INK, "Arial"),
                    ("Problem Statement Title –  AI-Powered IPsec VPN Protocol Analyzer and Security Assessment Framework", 16, True, INK, "Arial"),
                    ("Theme –  Cybersecurity  (NTRO — Blockchain & Cybersecurity)", 16, True, INK, "Arial"),
                    ("PS Category –  Software", 16, True, INK, "Arial"),
                    (f"Team ID –  {TEAM_ID}", 16, True, INK, "Arial"),
                    (f"Team Name (Registered on portal) –  {TEAM}", 16, True, INK, "Arial"),
                ],
            )


def fill_content_title(slide, title):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name.startswith("Title"):
            fill_tf(sh, [(title, 28, True, DK, "Times New Roman")])


def slide2(slide):
    fill_content_title(slide, "IDEA TITLE  ·  VPN SENTINEL")
    set_oval_team(slide)
    prepare_dashboard_crop()
    y0 = Inches(1.12)
    card(
        slide,
        Inches(0.28),
        y0,
        Inches(6.35),
        Inches(1.72),
        "PROBLEM",
        "IPsec VPN security depends on **cryptographic choices**, key exchange, and configuration that most tools never score. **Weak ciphers**, deprecated **IKEv1**, and small **DH groups** stay hidden unless an expert reads **Wireshark** by hand.",
        DK,
    )
    card(
        slide,
        Inches(6.75),
        y0,
        Inches(6.3),
        Inches(1.72),
        "OUR IDEA",
        "**VPN Sentinel** parses **IKE/ESP** from a PCAP, reads **IKE version**, cipher suite, and **DH group** when visible, then returns findings plus a **risk score**. Encrypted ESP payloads are **never decrypted**.",
        MID,
    )
    # Real dashboard (cropped) — full width, Helix-style tags on corners
    pic_l, pic_t, pic_w, pic_h = Inches(0.28), Inches(2.90), Inches(12.75), Inches(2.28)
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pic_l, pic_t, pic_w, pic_h)
    frame.adjustments[0] = 0.02
    frame.fill.solid()
    frame.fill.fore_color.rgb = DK
    frame.line.color.rgb = DK
    slide.shapes.add_picture(str(HERO), pic_l + Inches(0.04), pic_t + Inches(0.04), pic_w - Inches(0.08), pic_h - Inches(0.08))
    # Tags sit in a strip under the screenshot so the live UI is not covered
    tag_box(slide, Inches(0.28), Inches(5.22), Inches(3.05), Inches(0.34), "Live dashboard", BLUE)
    tag_box(slide, Inches(3.45), Inches(5.22), Inches(3.05), Inches(0.34), "Offline PCAP", DK)
    tag_box(slide, Inches(6.62), Inches(5.22), Inches(3.05), Inches(0.34), "Demo IKEv1 / IKEv2", MID)
    tag_box(slide, Inches(9.79), Inches(5.22), Inches(3.24), Inches(0.34), "No ESP decrypt", ORANGE)
    card(
        slide,
        Inches(0.28),
        Inches(5.60),
        Inches(6.35),
        Inches(1.20),
        "PROPOSED SOLUTION",
        "• Separate **RFC-correct IKEv1 and IKEv2** SA parsers — not one shared byte pattern.\n"
        "• Checks for **weak ciphers**, **weak DH**, **IKEv1**. **Planned:** IsolationForest on ESP metadata — **not** payload inspection.",
        BLUE,
    )
    card(
        slide,
        Inches(6.75),
        Inches(5.60),
        Inches(6.3),
        Inches(1.20),
        "INNOVATION / UNIQUENESS",
        "• **IKEv1 RFC 2409** vs **IKEv2 RFC 7296 §3.3** — shared hunters fail silently on IKEv2.\n"
        "• **Never decrypt** ESP. Parse + policy **risk score** from one offline PCAP.",
        ORANGE,
    )


def slide3(slide):
    fill_content_title(slide, "TECHNICAL APPROACH")
    set_oval_team(slide)
    card(
        slide,
        Inches(0.28),
        Inches(1.22),
        Inches(4.55),
        Inches(4.35),
        "SOFTWARE (NO HARDWARE APPLIANCE)",
        "• **FastAPI** — analysis API (PCAP upload)\n"
        "• **Scapy** — IKE / ESP / AH / UDP 500 / 4500\n"
        "• **parse_ikev1_sa** — RFC 2409 attributes\n"
        "• **parse_ikev2_sa** — RFC 7296 §3.3 transforms\n"
        "• Rule-based findings (IKEv1, weak ciphers, DH)\n"
        "• **Planned:** scikit-learn **IsolationForest** on ESP size/timing (**not** traffic-type classes)\n"
        "• **React (Vite)** — dashboard (screenshot on slide 2)\n"
        "• Input: **PCAP only** (no live sniff in this design)",
        DK,
    )
    # flowchart panel
    fl, ft, fw, fh = Inches(5.05), Inches(1.22), Inches(8.0), Inches(4.35)
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fl, ft, fw, fh)
    bg.adjustments[0] = 0.03
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.color.rgb = DK
    title = slide.shapes.add_textbox(fl + Inches(0.15), ft + Inches(0.06), fw - Inches(0.3), Inches(0.32))
    fill_tf(title, [("IMPLEMENTATION FLOW", 13, True, DK, "Arial")])

    y = Inches(1.62)
    arrow_box(slide, Inches(5.35), y, Inches(7.4), Inches(0.48), "PCAP UPLOAD  (offline capture)", BLUE)
    chevron(slide, Inches(8.85), y + Inches(0.50))
    y = Inches(2.28)
    arrow_box(slide, Inches(5.35), y, Inches(3.5), Inches(0.78), "IKEv1 PARSER\nRFC 2409  0x80 0x01 attrs", DK)
    arrow_box(slide, Inches(9.25), y, Inches(3.5), Inches(0.78), "IKEv2 PARSER\nRFC 7296  Type/ID transforms", MID)
    # branch label
    br = slide.shapes.add_textbox(Inches(5.35), Inches(3.08), Inches(7.4), Inches(0.28))
    fill_tf(br, [("Version nibble (**ISAKMP >> 4**) selects the branch — **not** the filename", 11, False, INK, "Arial")], align=PP_ALIGN.CENTER)
    chevron(slide, Inches(8.85), Inches(3.32))
    y = Inches(3.52)
    arrow_box(slide, Inches(5.35), y, Inches(3.5), Inches(0.72), "RULE ENGINE\nIKEv1 / weak cipher / DH", BLUE)
    arrow_box(slide, Inches(9.25), y, Inches(3.5), Inches(0.72), "PLANNED ML\nIsolationForest (ESP metadata)", ORANGE)
    chevron(slide, Inches(8.85), Inches(4.28))
    arrow_box(slide, Inches(5.35), Inches(4.50), Inches(7.4), Inches(0.48), "RISK SCORE  +  FINDINGS  →  DASHBOARD / PDF REPORT", DK)

    # bottom strip github + %
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.28), Inches(5.68), Inches(12.75), Inches(1.10))
    bar.adjustments[0] = 0.04
    bar.fill.solid()
    bar.fill.fore_color.rgb = PALE
    bar.line.color.rgb = BLUE
    tb = slide.shapes.add_textbox(Inches(0.45), Inches(5.78), Inches(12.4), Inches(0.90))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, "GitHub:  ", 13, True, DK, "Arial")
    r2 = p.add_run()
    set_run(r2, GITHUB, 13, False, BLUE, "Arial")
    r2.hyperlink.address = GITHUB
    p2 = tf.add_paragraph()
    r3 = p2.add_run()
    set_run(r3, f"Prototype completion:  {COMPLETION}", 13, True, INK, "Arial")
    p3 = tf.add_paragraph()
    r4 = p3.add_run()
    set_run(r4, "Count finished work only. Do not count unbuilt IsolationForest, testbed, or traffic-type AI.", 11, False, INK, "Arial")


def slide4(slide):
    fill_content_title(slide, "FEASIBILITY AND VIABILITY")
    set_oval_team(slide)
    card(
        slide,
        Inches(0.28),
        Inches(1.18),
        Inches(6.35),
        Inches(2.55),
        "FEASIBILITY",
        "• **Offline PCAP** analysis avoids live-capture complexity and legal/ops risk.\n"
        "• Parsers checked against real **IKEv1/IKEv2** captures using **Wireshark** as ground truth.\n"
        "• Findings follow published **NIST SP 800-77r1**-style IPsec guidance, not invented thresholds.",
        DK,
    )
    card(
        slide,
        Inches(6.75),
        Inches(1.18),
        Inches(6.3),
        Inches(2.55),
        "COMMERCIAL FEASIBILITY",
        "• Enterprises still run **IPsec**; automated first-pass audits reduce wait for a packet expert.\n"
        "• **MVP runs locally** (Python + React) — no appliance purchase for a demo.\n"
        "• Market-size figure only if sourced — **not invented**.",
        MID,
    )
    card(
        slide,
        Inches(0.28),
        Inches(3.88),
        Inches(6.35),
        Inches(2.85),
        "CHALLENGES AND RISKS",
        "• **IKEv1** and **IKEv2** need separate parsing (different RFCs). A shared **0x80 0x01** search is wrong on IKEv2 and fails silently.\n"
        "• ESP payloads are **encrypted** — no honest “this is YouTube” without a labelled dataset.\n"
        "• Tunnel/transport, **PFS**, and replay are often in encrypted Child SA — must stay **not detected** if absent.",
        RGBColor(0xC0, 0x50, 0x4D),
    )
    card(
        slide,
        Inches(6.75),
        Inches(3.88),
        Inches(6.3),
        Inches(2.85),
        "STRATEGY",
        "• **Planned:** IsolationForest on ESP flow metadata — **not** multi-class traffic-type claims.\n"
        "• **Testbed** and labelled web/video/voice classification = **future work**.\n"
        "• Findings stay explainable: **rule IDs + packet evidence**; ML is a separate finding when wired.",
        RGBColor(0x9B, 0xBB, 0x59),
    )


def slide5(slide):
    fill_content_title(slide, "IMPACT AND BENEFITS")
    set_oval_team(slide)
    card(
        slide,
        Inches(0.28),
        Inches(1.18),
        Inches(6.35),
        Inches(3.55),
        "DIRECT IMPACT ON TARGET USERS",
        "• Security teams: automated first-pass VPN audit instead of only manual **Wireshark**.\n"
        "• Admins: **IKEv1**, **DES/3DES**, and **weak DH** called out with a recommendation.\n"
        "• Reduces expert-hours on the handshake and advertised transforms.\n"
        "• Does **not** replace a full NTRO testbed or **decrypt** user traffic.",
        DK,
    )
    card(
        slide,
        Inches(6.75),
        Inches(1.18),
        Inches(6.3),
        Inches(3.55),
        "STRATEGIC IMPACT",
        "• Supports correctly configured IPsec (**IKEv2**, modern ciphers, adequate **DH**).\n"
        "• Makes **“VPN is on”** distinguishable from **“the handshake is actually strong.”**\n"
        "• Aligns with **PS 26160** assessment + report; testbed and inner-traffic class remain the expansion path.",
        BLUE,
    )
    card(
        slide,
        Inches(0.28),
        Inches(4.88),
        Inches(12.75),
        Inches(1.85),
        "ECONOMIC / SOCIAL BENEFITS",
        "• Social: fewer unnoticed **weak VPN handshakes** on captures the operator is allowed to inspect.\n"
        "• Economic: no fabricated crore figure — software-only student/lab cost.\n"
        "• Environmental: **PCAP + optional containers**, not extra dedicated hardware for the prototype.",
        ORANGE,
    )


def slide6(slide):
    fill_content_title(slide, "RESEARCH  AND REFERENCES")
    set_oval_team(slide)
    cells = [
        (0.28, 1.18, DK, "GAP & PROBLEM IDENTIFICATION",
         "Misconfigured IPsec (**IKEv1**, **weak DH**, legacy ciphers) is hard to see without packet expertise. **PS 26160** asks for automated identification + assessment; generic “AI VPN” tools often skip **RFC-correct** SA layout."),
        (6.75, 1.18, BLUE, "LITERATURE SURVEY",
         "**NIST SP 800-77r1** — IPsec configuration guidance.\n**Cremers, ESORICS 2011** — IKEv1 vs IKEv2 security.\n**Draper-Gil et al., ICISSP 2016** — time-related features for encrypted/VPN traffic (metadata-only ESP)."),
        (0.28, 3.05, MID, "TECHNOLOGY BENCHMARKING",
         "**ike-scan / iker** — enumerate IKE transforms from live probes.\n**VPN Sentinel** — same field goal, **offline from PCAP**, separate v1/v2 parsers.\n**Wireshark** — ground-truth decode, not replaced."),
        (6.75, 3.05, ORANGE, "ECONOMIC LANDSCAPE",
         "Demand is cheaper **first-pass audits**, not a new VPN vendor.\nOpen stack (**Python, Scapy, FastAPI**) keeps prototype cost at student/lab scale."),
        (0.28, 4.92, RGBColor(0x9B, 0xBB, 0x59), "FIELD TESTS / VALIDATION",
         "**Wireshark-verified** IKEv2 sample: **ENCR_AES_CBC (12)** key **256**, PRF SHA2-512, INTEG SHA2-512-256, **DH group 20**.\nIKEv1 sample: **AES-CBC**, key 256, DH 20, **PSK** attributes."),
        (6.75, 4.92, RGBColor(0x80, 0x64, 0xA2), "POLICY / ECOSYSTEM",
         "**RFC 9395** — IKEv1 deprecation (a finding, not a slogan).\nOperator-owned captures only; **no decrypt**; no scanning third-party networks.\nSIH software prototype: **local demo + docs + labelled samples**."),
    ]
    for x, y, c, title, body in cells:
        card(slide, Inches(x), Inches(y), Inches(6.28), Inches(1.75), title, body, c)


def main():
    prepare_dashboard_crop()
    shutil.copy(TEMPLATE, OUT)
    prs = Presentation(str(OUT))
    slides = list(prs.slides)
    fill_slide1(slides[0])
    delete_instruction_boxes(slides[1])
    slide2(slides[1])
    delete_instruction_boxes(slides[2])
    slide3(slides[2])
    delete_instruction_boxes(slides[3])
    slide4(slides[3])
    delete_instruction_boxes(slides[4])
    slide5(slides[4])
    delete_instruction_boxes(slides[5])
    slide6(slides[5])
    delete_last_slide(prs)
    prs.save(str(OUT))
    shutil.copy(OUT, Path("/workspace/docs/VPN_Sentinel_SIH2026_IDEA.pptx"))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
